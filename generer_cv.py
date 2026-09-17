# -*- coding: utf-8 -*-
"""Génère le CV ATS (1 page, Word) à partir des infos du stage Alumni CRM."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

NAVY = RGBColor(0x15, 0x20, 0x42)
GRAY = RGBColor(0x60, 0x6A, 0x7E)

doc = Document()

section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.0)
section.bottom_margin = Cm(1.0)
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)
style.paragraph_format.space_after = Pt(2)
style.paragraph_format.line_spacing = 1.05


def heading(text, size=13):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = NAVY
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr")
    if pBdr is None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "15 20 42")
    pBdr.append(bottom)
    return p


def bullet(label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    r = p.add_run("• ")
    r.font.color.rgb = NAVY
    if label:
        r2 = p.add_run(label)
        r2.bold = True
        if text:
            p.add_run(" — " + text)
    else:
        p.add_run(text)
    return p


# ── En-tête ─────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("RAFIK DJEMADI")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = NAVY

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Étudiant Pré-MSc Informatique & Management — IONIS-STM")
r.font.size = Pt(12)
r.font.color.rgb = GRAY

contacts = [
    "Email : [tonemail@exemple.com]",
    "Téléphone : [06 XX XX XX XX]",
    "Localisation : [Ville]",
    "LinkedIn : linkedin.com/in/rafik-djemadi",
    "GitHub : github.com/[ton-compte]",
]
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, c in enumerate(contacts):
    if i:
        r = p.add_run("    |    ")
        r.font.color.rgb = GRAY
    r = p.add_run(c)
    r.font.size = Pt(9)
    r.font.color.rgb = GRAY

# ── Profil ──────────────────────────────────────────────────────
heading("Profil")
bullet("", "Étudiant Pré-MSc en reconversion vers la double compétence Tech & Management, spécialisé en "
           "développement web full-stack. Auteur d'un projet complet de conception et développement : un CRM "
           "de suivi des anciens étudiants (backend FastAPI, frontend React, base PostgreSQL), mené en solo "
           "avec conformité RGPD, indicateurs d'insertion et livrables Management — du cadrage à la livraison.")

# ── Formation ───────────────────────────────────────────────────
heading("Formation")
bullet("", "Pré-MSc Informatique & Management — IONIS-STM, Paris (2026) — première école française dédiée "
           "à la double compétence Tech & Management.")
bullet("", "Programme en reconversion et poursuite d'études après un premier diplôme ; filières : "
           "développement, cybersécurité, data, management, marketing digital.")

# ── Expérience / Réalisation principale ─────────────────────────
heading("Expérience — Stage de mise en situation professionnelle")
bullet("Conception & développement d'Alumni CRM (application de gestion des alumni)", "")
bullet("", "Stage de substitution proposé par IONIS-STM — 2026 (solo, autonomie totale sur les choix techniques).")
bullet("", "Objectif : suivre le parcours complet de l'étudiant, de l'inscription à sa vie professionnelle, et "
           "valoriser le réseau des anciens, en restant conforme au RGPD.")
bullet("", "Backend : API REST FastAPI (Python), 84 endpoints, documentation Swagger automatique, 17 migrations "
           "PostgreSQL versionnées, 14 tables normalisées en 5 domaines.")
bullet("", "Frontend : application web React + Vite, 14 routes, deux espaces séparés (admin / alumni), "
           "thème clair / sombre.")
bullet("", "Sécurité : clé API admin, sessions JWT par rôle, code OTP à 6 chiffres, vérification des droits "
           "d'accès, gestion des erreurs.")
bullet("", "RGPD : 4 types de consentement tracés, export et suppression via workflow verrouillé, décision "
           "anonymiser vs supprimer, purge différée, cartographie des données personnelles.")
bullet("", "Indicateurs : 8 indicateurs d'insertion automatisés (taux d'emploi, salaire, secteur…), "
           "tableau de bord KPI admin, import / export Excel.")
bullet("", "Livrables Management : cartographie des données, charte RGPD, analyse des indicateurs d'insertion, "
           "stratégie de mise à jour des données, guide des processus d'animation du réseau.")
bullet("", "Audits : cohérence base / API et revue de sécurité (correctifs versionnés).")

# ── Compétences ─────────────────────────────────────────────────
heading("Compétences")
bullet("Développement full-stack", "Python, FastAPI, REST / JSON, React, Vite, JavaScript, HTML/CSS.")
bullet("Base de données", "PostgreSQL, SQL, modélisation MCD / MLD (Looping), migrations versionnées.")
bullet("Sécurité & conformité", "JWT, OTP, contrôle d'accès, RGPD (consentements, extraction, effacement).")
bullet("Management", "Analyse des indicateurs, processus d'animation de réseau, documentation technique.")
bullet("Outils", "Git, Excel (import / export), fpdf2 / python-docx (génération de documents).")

# ── Langues ─────────────────────────────────────────────────────
heading("Langues")
bullet("", "Français : courant    |    Anglais : [niveau — ex. professionnel]")

# ── Note ATS ────────────────────────────────────────────────────
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
r = p.add_run("Note : remplace les éléments entre crochets [ ] (contacts, GitHub, niveau d'anglais) "
              "avant envoi. Structuré en une colonne et sans images : compatible ATS.")
r.font.size = Pt(8)
r.font.color.rgb = GRAY
r.italic = True

OUT = r"C:\Users\PC\OneDrive\Desktop\stage\CV - Rafik Djemadi.docx"
doc.save(OUT)
print("OK :", OUT)